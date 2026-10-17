"""
Script tao file Word mau (.docx) theo chuan bo the dinh nghia trong AGENT_PLAN.md
Tu de ETS TOEIC Reading Test 09 (Reading v09.pdf).
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor

def create_sample_docx(output_path, is_valid=True):
    doc = docx.Document()
    
    # Custom styles or formatting
    title_p = doc.add_paragraph("[EXAM_TITLE: ETS TOEIC Reading Test 09]")
    doc.add_paragraph("[EXAM_DURATION: 75]")
    doc.add_paragraph("")
    
    # ==========================================
    # PART 5
    # ==========================================
    doc.add_paragraph("[PART: 5]")
    
    part5_data = [
        (101, "Every batch of sauce at Generita's Bistro is processed meticulously by ------- expert chefs.",
         "they", "their", "them", "themselves", "B", "Cần tính từ sở hữu 'their' đứng trước cụm danh từ 'expert chefs'."),
        (102, "Computerization of medical records ------- increases a physician's ability to diagnose and treat patients.",
         "great", "greatly", "greatness", "greatest", "B", "Cần trạng từ 'greatly' bổ nghĩa cho động từ 'increases'."),
        (103, "At Rojelle's Fine Dining, we use the freshest ------- available to make our salads.",
         "applications", "subjects", "ingredients", "factors", "C", "Dựa theo ngữ cảnh nhà hàng, 'ingredients' (nguyên liệu) tươi ngon để làm salad."),
        (104, "Professor Benguigui will present his paper ------- the natural history conference tomorrow.",
         "by", "at", "of", "on", "B", "Giới từ 'at' đi với sự kiện/hội nghị: 'at the conference'."),
        (105, "The chefs are still waiting for a ------- date for the new convection oven that was ordered two weeks ago.",
         "shipper", "ships", "shipments", "shipping", "D", "Cụm danh từ: 'shipping date' nghĩa là ngày giao hàng."),
        (106, "Mr. Kang works ------- with our internal team members as well as various regional sales representatives.",
         "mildly", "nearly", "closely", "narrowly", "C", "Cụm 'work closely with' nghĩa là làm việc chặt chẽ/gần gũi với ai."),
        (107, "Fashion designer Hye-Ja Pak knows ------- to update her line in response to changing tastes.",
         "and", "when", "need", "for", "B", "Cấu trúc từ để hỏi + to V: 'when to update' (khi nào cần cập nhật dòng sản phẩm)."),
        (108, "The samba class was so well ------- that the Yulara Dance School decided to make the course a permanent offering.",
         "attended", "educated", "gathered", "protected", "A", "'well attended' có nghĩa là được đông đảo học viên tham dự."),
        (109, "The factory will be situated away ------- the city's residential area to reduce complaints about noise and emissions.",
         "from", "about", "with", "out", "A", "Cụm 'away from' nghĩa là cách xa khỏi khu dân cư."),
        (110, "Ms. Kuramoto selected the most ------- mailing option available.",
         "economical", "economy", "economize", "economized", "A", "Cần tính từ 'economical' (tiết kiệm) bổ nghĩa cho danh từ 'mailing option'."),
        (111, "Companies without information technology specialists can ------- on Vyber Software Advisers for assistance with online services.",
         "reliable", "reliably", "rely", "relying", "C", "Sau trợ động từ 'can' là động từ nguyên mẫu 'rely' (rely on: dựa vào, trông cậy vào)."),
        (112, "Next year, our team will have a new task, ------- is to review design portfolios.",
         "although", "which", "after", "because", "B", "Đại từ quan hệ 'which' thay thế cho mệnh đề hoặc danh từ 'a new task'."),
        (113, "Cook the meat for 30 minutes to ensure ------- readiness to be eaten.",
         "both", "this", "its", "that", "C", "Cần tính từ sở hữu 'its' bổ nghĩa cho danh từ 'readiness'."),
        (114, "Amoxitron's research team will hire ------- interns to assist with laboratory duties.",
         "given", "several", "whole", "natural", "B", "'several' (một vài/nhiều) đi với danh từ số nhiều đếm được 'interns'."),
        (115, "------- the new acai juice blend has proved so popular, we should move quickly to increase our production volume.",
         "If", "Whether", "Since", "Unless", "C", "'Since' (bởi vì) giải thích lý do món nước uống được ưa chuộng nên cần tăng sản lượng."),
        (116, "New emissions standards have forced Rider Auto to modify the process of engine -------.",
         "construction", "constructed", "constructive", "construct", "A", "Sau giới từ 'of' cần danh từ 'construction' (quá trình chế tạo động cơ)."),
        (117, "Maki Kayano's book offers techniques for ------- business tasks with speed and precision.",
         "executing", "equipping", "returning", "involving", "A", "'executing business tasks' nghĩa là thực hiện/thi hành các nhiệm vụ kinh doanh."),
        (118, "To control costs, updated credit card readers will be installed in branch stores -------.",
         "gradual", "gradually", "more gradual", "most gradual", "B", "Cần trạng từ 'gradually' (dần dần) bổ nghĩa cho cụm động từ 'will be installed'."),
        (119, "In this quarter, the Mantel Beverage Company is expecting sales ------- £160,000 and £180,000.",
         "without", "among", "throughout", "between", "D", "Cấu trúc 'between A and B' (ở mức dao động giữa £160,000 và £180,000)."),
        (120, "Because Mount Akoyola is so challenging for climbers, ------- have reached its peak.",
         "any", "either", "other", "few", "D", "'few' (rất ít người) chỉ số lượng người chinh phục được đỉnh núi vì núi quá hiểm trở."),
        (121, "King Street Bridge will be closed in the month of September ------- repair work.",
         "because of", "so that", "as if", "rather than", "A", "'because of' đi với cụm danh từ 'repair work' (do công tác sửa chữa)."),
        (122, "Ms. Taniguchi's supervisor commended her for negotiating ------- with Furuyama Corporation.",
         "effective", "effecting", "effected", "effectively", "D", "Cần trạng từ 'effectively' bổ nghĩa cho động từ 'negotiating' (đàm phán hiệu quả)."),
        (123, "Staff members may reserve the conference room ------- they need it.",
         "somewhere", "whatever", "everything", "anytime", "D", "'anytime' (bất kỳ lúc nào) đóng vai trò liên từ chỉ thời gian."),
        (124, "The public relations director must have a high level of ------- in English and Spanish.",
         "proficiency", "advancement", "routine", "strength", "A", "'proficiency in' nghĩa là sự thành thạo/thông thạo một ngôn ngữ."),
        (125, "Patterson Products ------- seeks innovative ways of meeting changing consumer demand.",
         "dually", "favorably", "continually", "generically", "C", "Trạng từ 'continually' (liên tục/không ngừng) tìm kiếm các giải pháp sáng tạo."),
        (126, "Mr. Volante is working at home tomorrow so he ------- the technical report without any distractions.",
         "can finish", "would finish", "finished", "has been finishing", "A", "Mệnh đề chỉ mục đích 'so he can finish' (để anh ấy có thể hoàn thành báo cáo)."),
        (127, "All of the billing procedures remain the same, ------- new payment codes need to be inserted into the invoice documents.",
         "in order that", "during", "across from", "except that", "D", "'except that' (ngoại trừ việc) chỉ điểm khác biệt duy nhất."),
        (128, "The lightweight design of the new sedan is ------- because it can cause the vehicle to slide on icy roads.",
         "possible", "mechanical", "questionable", "multiple", "C", "'questionable' (đáng nghi ngại/đáng quan ngại) vì dễ làm xe trượt trên đường băng."),
        (129, "When the bank president retired, the common ------- was that the vice president would take over.",
         "assume", "assumption", "assuming", "assumable", "B", "Sau tính từ 'common' cần danh từ 'assumption' (nhận định/giả định chung)."),
        (130, "West Bengali Airlines ------- fees for oversized items that are still within weight limitations.",
         "invites", "cooperates", "transports", "waives", "D", "Động từ 'waives' (miễn trừ phí) đối với các kiện hàng quá khổ nhưng trong giới hạn cân nặng.")
    ]
    
    for q in part5_data:
        q_num, text, opt_a, opt_b, opt_c, opt_d, ans, hint = q
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: TRUE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        doc.add_paragraph(f"[A] {opt_a}")
        doc.add_paragraph(f"[B] {opt_b}")
        doc.add_paragraph(f"[C] {opt_c}")
        # Test case: intentionally omit [D] if invalid test file
        if not is_valid and q_num == 105:
            pass # omit [D]
        else:
            doc.add_paragraph(f"[D] {opt_d}")
        doc.add_paragraph(f"[ANS: {ans}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[Q_END]")
        doc.add_paragraph("")
        
    # ==========================================
    # PART 6
    # ==========================================
    doc.add_paragraph("[PART: 6]")
    
    part6_passages = [
        # Passage 1 (131-134)
        {
            "text": """Questions 131-134 refer to the following e-mail.
To: pmendoza@factmail.co
From: recruiting@analystsassoc.org
Date: May 2
Subject: Joining ASA
Dear Mr. Mendoza,
Thank you for expressing your interest in our organization during the recent ---[131]---. It was a pleasure meeting you at the ASA booth during the Weber Information Systems Convention. As you may recall from our conversation, we discussed how ASA membership ---[132]--- your career through networking opportunities as well as the industry insights offered in our monthly newsletter. The normal fee for members is just $120 a year; however, we are currently offering a new-member discount. ---[133]---.
If you are still interested in joining, please reply to ---[134]--- with your mailing address. I will then forward you an application packet.
Sincerely,
Ashlee Loren, President
Association of Systems Analysts""",
            "questions": [
                (131, "Chọn từ thích hợp điền vào chỗ trống [131]:", "election", "broadcast", "conference", "performance", "C", "Câu sau có nhắc tới 'Weber Information Systems Convention' đồng nghĩa với 'conference'."),
                (132, "Chọn dạng động từ thích hợp điền vào chỗ trống [132]:", "can benefit", "is benefiting", "has benefited", "will have benefited", "A", "Hội viên ASA 'can benefit' (có thể mang lại lợi ích cho) sự nghiệp của bạn."),
                (133, "Chọn câu thích hợp điền vào chỗ trống [133]:", "Discounts on rental cars are included in the annual membership.", "Our jobs board is quite comprehensive.", "We are the first organization of our kind.", "This month you can join for just half the usual rate.", "D", "Phía trước đang nhắc đến 'new-member discount' nên câu giải thích chi tiết mức giá giảm một nửa là phù hợp nhất."),
                (134, "Chọn đại từ thích hợp điền vào chỗ trống [134]:", "it", "me", "them", "anyone", "B", "'reply to me' (hãy phản hồi cho tôi kèm địa chỉ thư tín của bạn).")
            ]
        },
        # Passage 2 (135-138)
        {
            "text": """Questions 135-138 refer to the following e-mail.
To: Sam Heinz
From: Northways Professional Development
Sent: April 20
Subject: Workshop 4/28-4/30
Dear Workshop Participants,
We look forward to seeing you at the digital storytelling workshop. On the first day, when you enter the Albin College campus, attendants will ---[135]--- you to lot 43 and the Toteman Building. We will begin each day in conference room 9. Coffee, tea, snacks, and fruit ---[136]--- in the mornings. Lunch will be sandwiches and salads from Black Horse Restaurant. ---[137]---.
In the afternoons, we will be working in the computer lab. We suggest that you assemble some images that you would like to use for your project. It will be ---[138]--- for you to have them saved on a digital storage device beforehand. If you have any questions, please e-mail us.
Very Best,
Gina Kapuski""",
            "questions": [
                (135, "Chọn từ thích hợp điền vào chỗ trống [135]:", "offer", "direct", "pass", "instruct", "B", "Nhân viên hướng dẫn sẽ 'direct you' (hướng dẫn/chỉ đường cho bạn đến bãi xe 43)."),
                (136, "Chọn thì thích hợp điền vào chỗ trống [136]:", "will be provided", "were provided", "providers", "are providing", "A", "Thì tương lai bị động 'will be provided' (đồ uống, thức ăn nhẹ sẽ được cung cấp)."),
                (137, "Chọn câu thích hợp điền vào chỗ trống [137]:", "Some prior experience working with digital files is assumed.", "Please do not bring any copyrighted material to the workshop.", "Please let us know if you have any dietary restrictions.", "There is a one-time parking charge of fifteen dollars.", "C", "Phía trước vừa đề cập đến bữa trưa bánh mì và salad nên câu nhắc về dị ứng thực phẩm là chuẩn xác nhất."),
                (138, "Chọn tính từ thích hợp điền vào chỗ trống [138]:", "useful", "surprised", "difficult", "amazing", "A", "'It will be useful' (Sẽ rất hữu ích nếu bạn lưu ảnh sẵn vào thiết bị lưu trữ số).")
            ]
        },
        # Passage 3 (139-142)
        {
            "text": """Questions 139-142 refer to the following article.
Venley Foods Responsive to Changing Consumers
Who cares where the tomatoes in your salad actually came from? ---[139]---, an increasing number of people do, according to a study conducted by the Consumer Group. In fact, the study shows that many consumers would pay an average of 10 percent more when given the exact source of a fresh food product. ---[140]---. Some grocery stores, such as Venley Foods in Boston, have taken advantage of the trend and used it to implement ---[141]--- branding and marketing. "If we can tell a story about our product," says Venley Foods CEO Minji Kim, "then we've added ---[142]--- in the minds of consumers.\"""",
            "questions": [
                (139, "Chọn từ nối thích hợp điền vào chỗ trống [139]:", "Carefully", "Apparently", "Formerly", "Rarely", "B", "'Apparently' (Rõ ràng là/Hóa ra là) người tiêu dùng ngày càng quan tâm đến nguồn gốc thực phẩm."),
                (140, "Chọn câu thích hợp điền vào chỗ trống [140]:", "Fresh food can be refrigerated for up to two days.", "Many grocery stores have been extending their hours.", "Most studies are published in consumer magazines.", "The number increases to 20 percent in large cities.", "D", "Phía trước nhắc tới 'pay an average of 10 percent more', câu bổ sung thông tin tỉ lệ này lên đến 20% ở các đô thị lớn."),
                (141, "Chọn từ thích hợp điền vào chỗ trống [141]:", "smarts", "smartly", "smarter", "smartness", "C", "Tính từ so sánh 'smarter branding and marketing' (xây dựng thương hiệu và tiếp thị thông minh hơn)."),
                (142, "Chọn danh từ thích hợp điền vào chỗ trống [142]:", "value", "time", "obstacles", "bonus", "A", "'added value' (gia tăng giá trị) trong tâm trí người tiêu dùng.")
            ]
        },
        # Passage 4 (143-146)
        {
            "text": """Questions 143-146 refer to the following article.
Nylobe, Inc., Announces Newest Development Project
SEATTLE (May 1)- Researchers at technology firm Nylobe, Inc., are working to develop a sensor capable of detecting corrosion caused by environmental exposure. Corrosion is a major contributor to ---[143]--- losses in the aircraft industry each year. "This will be a major ---[144]--- for commercial airline fleets," says Mel Laveau, Nylobe's CEO. "The sensor will decrease both labor and maintenance costs without being too expensive."
According to Ms. Laveau, the sensor will work by detecting corrosion in its early stages, when the problem can be corrected simply by removing the corroded material. ---[145]---. In the structure of large aircraft, some critical joints can be particularly susceptible to corrosion. ---[146]---, the sensor can be used to inspect these areas and then target the most likely areas of concern.""",
            "questions": [
                (143, "Chọn từ thích hợp điền vào chỗ trống [143]:", "financially", "financed", "financial", "finances", "C", "Cần tính từ 'financial' bổ nghĩa cho danh từ 'losses' (những tổn thất tài chính)."),
                (144, "Chọn danh từ thích hợp điền vào chỗ trống [144]:", "balance", "examination", "expectation", "asset", "D", "'a major asset' (một tài sản/lợi thế lớn đối với các đội tàu bay thương mại)."),
                (145, "Chọn câu thích hợp điền vào chỗ trống [145]:", "This will reduce the need for making expensive structural repairs.", "The parts have all been replaced with higher quality materials.", "The next stage of the project involves scanning the affected areas.", "Its style and sleek design made it popular with the public.", "A", "Phát hiện ăn mòn từ sớm sẽ giúp giảm bớt việc phải sửa chữa kết cấu tốn kém."),
                (146, "Chọn trạng từ liên kết thích hợp điền vào chỗ trống [146]:", "Meanwhile", "Similarly", "Otherwise", "Fortunately", "D", "'Fortunately' (Thật may mắn là cảm biến có thể kiểm tra những khu vực hiểm yếu này).")
            ]
        }
    ]
    
    for p in part6_passages:
        doc.add_paragraph("[PASSAGE_START]")
        doc.add_paragraph(p["text"])
        doc.add_paragraph("[PASSAGE_END]")
        doc.add_paragraph("")
        
        for q in p["questions"]:
            q_num, text, opt_a, opt_b, opt_c, opt_d, ans, hint = q
            doc.add_paragraph(f"[Q: {q_num}] {text}")
            doc.add_paragraph("[SHUFFLE_Q: FALSE]")
            doc.add_paragraph("[SHUFFLE_A: TRUE]")
            doc.add_paragraph(f"[A] {opt_a}")
            doc.add_paragraph(f"[B] {opt_b}")
            doc.add_paragraph(f"[C] {opt_c}")
            doc.add_paragraph(f"[D] {opt_d}")
            doc.add_paragraph(f"[ANS: {ans}]")
            doc.add_paragraph(f"[HINT: {hint}]")
            doc.add_paragraph("[Q_END]")
            doc.add_paragraph("")
            
    # ==========================================
    # PART 7
    # ==========================================
    doc.add_paragraph("[PART: 7]")
    
    part7_passages = [
        # 147-148 Invoice
        {
            "text": """Questions 147-148 refer to the following invoice.
Invoice 3987 - Shawqi Office Services, Dubai, UAE
18 October | Report requested in English
Customer: Jenkins Press, P.O. Box 2291, Dubai, UAE
Service Description:
On 14 October, replaced bulb and repaired paper tray on copier per call received on 12 October. Replaced copy ink in two machines. Performed routine yearly maintenance on five copiers per existing service contract.
Labor cost: AED 330.00
Paper tray: AED 50.00
Bulb: AED 30.00
Copy ink: AED 220.00
Total: AED 630.00
Total amount must be received by 31 October.
Thank you for your business!""",
            "questions": [
                (147, "What is indicated about Jenkins Press?",
                 "It has several offices around the world.", "Its copiers get checked every year.", "Its office equipment is outdated.", "It is a new customer of Shawqi Office Services.",
                 "B", "Hóa đơn ghi: 'Performed routine yearly maintenance on five copiers per existing service contract.'"),
                (148, "When is payment due?",
                 "October 12", "October 14", "October 18", "October 31",
                 "D", "Hóa đơn ghi rõ hạn thanh toán: 'Total amount must be received by 31 October.'")
            ]
        },
        # 149-151 Advertisement
        {
            "text": """Questions 149-151 refer to the following advertisement.
Harbour View Apartment- Porthmadog, Wales
This one-bedroom apartment is perfect for a holiday escape! Located in a quiet area, it boasts a patio with a lovely view of the harbour. Recently renovated, the unit includes an eat-in kitchen with stove, refrigerator, microwave, and coffeepot; bathroom with walk-in shower; and a living room with a large-screen TV. Other amenities of the property include:
• Short distance to restaurants and shops
• Five-minute walk to the beach
• Public gardens and historic sites within a 20-minute drive
• Heat and electricity included
• Towels and bed linens provided on-site
• Daily cleaning service available (extra fee)
• Wireless Internet access (extra fee)
Reserve this lovely gem now! Signing a contract by March 30 will reduce the rental cost by 10 percent. To sign a contract, contact Dylan Barrett at dbarrett@telarentals.co.uk.""",
            "questions": [
                (149, "Who would the advertisement most likely interest?",
                 "Business travelers", "Residents of Porthmadog", "Property investors", "Short-term vacationers",
                 "D", "Bài quảng cáo mở đầu: 'perfect for a holiday escape' (khách nghỉ dưỡng ngắn hạn)."),
                (150, "What is indicated about the rental fee?",
                 "It includes tours of historic places.", "It does not cover all of the apartment's features.", "It includes vouchers to use at local restaurants.", "It requires a minimum 30 percent deposit in advance.",
                 "B", "Dọn dẹp hằng ngày và wifi bị tính phí thêm ('extra fee'), nghĩa là giá thuê cơ bản không bao gồm tất cả tiện ích."),
                (151, "Why should an individual contact Mr. Barrett by March 30?",
                 "To get a discount on rent", "To schedule apartment renovations", "To rent the last available property", "To sell the property before the end of the season",
                 "A", "Bài viết ghi: 'Signing a contract by March 30 will reduce the rental cost by 10 percent.'")
            ]
        },
        # 152-153 Text message chain
        {
            "text": """Questions 152-153 refer to the following text-message chain.
Jason Salter (1:45 P.M.)
Liz, the meeting is starting in fifteen minutes. Where are you?

Liz Ortiz (1:47 P.M.)
The train has been stopped on the tracks for a while. There seems to be some kind of problem. I still hope to make the meeting.

Jason Salter (1:50 P.M.)
Okay. I'll save you a seat.

Liz Ortiz (1:59 P.M.)
There's just been an announcement. There's a disabled train up ahead. It'll be awhile.

Jason Salter (2:00 P.M.)
Don't worry. If there are any questions for our department, I'll handle them.

Liz Ortiz (2:01 P.M.)
Thanks. I'll call you later.""",
            "questions": [
                (152, "What is suggested about Mr. Salter?",
                 "He is leading the meeting.", "He commutes to work by train.", "He has many questions for Ms. Ortiz.", "He works in the same department as Ms. Ortiz.",
                 "D", "Mr. Salter nhắn: 'If there are any questions for our department, I'll handle them.'"),
                (153, "At 1:59 P.M., what does Ms. Ortiz imply when she writes, 'It'll be awhile'?",
                 "The meeting is running late.", "She is still preparing her notes.", "She will likely miss the meeting.", "She has not boarded the train yet.",
                 "C", "Tàu hỏng phía trước mất nhiều thời gian nên cô ấy khó có thể đến kịp cuộc họp lúc 2:00 P.M.")
            ]
        },
        # 154-155 E-mail
        {
            "text": """Questions 154-155 refer to the following e-mail.
From: tbogosian@sardhasconvention.com
To: pradalaily@dmcv.com
Date: 22 March
Subject: Your reservation
Attachment: Prada party

Dear Ms. Prada,
Thank you for choosing Sardha's Banquet Hall to host your event on 30 August. Unfortunately, West Hall will be under renovation in August, and we will be unable to accommodate your group in that room. However, your party will fit comfortably in East Hall. This room features floor-to-ceiling windows and a full view of the river. The 15 percent deposit we have received will hold the reservation.
As indicated when you completed your request through our Web site, your total price of $1,600 includes a full-course meal that includes appetizers, soup or salad, main course, and dessert. Attached to this e-mail please find a copy of the menu presented to guests with options for appetizer and main course.

Tomas Bogosian, General Manager""",
            "questions": [
                (154, "What is one purpose of the e-mail?",
                 "To cancel a reservation", "To apologize for an error", "To indicate a room change", "To request an extra deposit",
                 "C", "Thông báo đổi phòng tiệc từ West Hall sang East Hall do sửa chữa."),
                (155, "What is suggested about Ms. Prada?",
                 "She has paid her bill in full.", "She eats at Sardha's regularly.", "She works in the food industry.", "She made the reservation online.",
                 "D", "Email có câu: 'when you completed your request through our Web site' (đặt chỗ trực tuyến).")
            ]
        },
        # 156-157 Memo
        {
            "text": """Questions 156-157 refer to the following memo.
To: All Employees
From: IT Supervisor
Subject: OS Update
Date: 27 May

All company computers in the Melbourne and Victoria offices require an update to the operating system. This update will prepare our computers for the new version of our accounting software, which should arrive in the first week of July. Technicians will install the new operating system beginning on 3 June. We expect the process to be completed around 16 June, providing time to work out any bugs with the operating system before the accounting software is uploaded.
Completing the entire process will require you to log in so that we can verify the system is working as anticipated. Therefore, if you plan to go on holiday during this period, please inform IT Support immediately of the dates you will be out so that we can plan an alternate date to accommodate your schedule.
If you have any questions, please contact IT Support staff at extension 48.""",
            "questions": [
                (156, "When is the installation of the operating system expected to be finished?",
                 "In the first week of June", "In the middle of June", "At the end of June", "In the first week of July",
                 "B", "Memo ghi: 'We expect the process to be completed around 16 June' (giữa tháng 6)."),
                (157, "Who is asked to contact IT Support?",
                 "Employees who do not need the accounting software", "Accountants already using the new operating system", "Staff who work outside the offices", "Employees who are taking time off",
                 "D", "Memo yêu cầu: 'if you plan to go on holiday during this period, please inform IT Support immediately'.")
            ]
        },
        # 158-160 Letter
        {
            "text": """Questions 158-160 refer to the following letter.
Clarke-Ellis Construction
#20 Murphy Industrial Park, St. Michael BB23028 | 1-246-555-0126
Full-service commercial contractor serving all of Barbados

4 June
Ida Gutierrez, Darling Cove Inn, Mango Drive, Folkestone BB24017
Dear Ms. Gutierrez,
Thank you for contacting Clarke-Ellis Construction for your roofing project. ---[1]---. After inspecting the property, I have confirmed that the inn's main roof is in good condition and requires no repairs at this time. However, the roof of the inn's porch appears to be at least twenty years old and is worn beyond repair. ---[2]---. Clarke-Ellis Construction can remove and dispose of the existing porch roof and install a new one. The replacement will be comparable to the quality, style, and colour of the inn's main roof. We will use only commercial-grade leak barriers, insulation, and shingles manufactured by West Indies Weatherproofing, Inc. The estimated total cost, inclusive of labour and materials, is $3,260. ---[3]---. Additional fees would apply should you want us to make other improvements, such as painting the porch or replacing porch screens.
Please call me at the number above to discuss scheduling or any questions you have about the work. ---[4]---. I hope to hear from you soon.

Sincerely,
Grayson Clarke, Co-owner, Clarke-Ellis Construction""",
            "questions": [
                (158, "Why did Mr. Clarke write the letter?",
                 "To provide a work proposal", "To ask for a project extension", "To request an inspection report", "To submit a revised cost estimate",
                 "A", "Thư gửi báo giá và đề xuất phương án thi công mái lợp ('work proposal')."),
                (159, "What is indicated about West Indies Weatherproofing, Inc.?",
                 "It is installing a new porch at Darling Cove Inn.", "It provides painting services.", "It produces roofing materials.", "It is owned by Clarke-Ellis Construction.",
                 "C", "Thư đề cập các vật liệu như leak barriers, insulation, and shingles được sản xuất bởi công ty này."),
                (160, "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong? 'It must be replaced.'",
                 "[1]", "[2]", "[3]", "[4]",
                 "B", "Đứng sau câu mô tả mái hiên bị mục rách không thể sửa chữa ('worn beyond repair -> It must be replaced').")
            ]
        },
        # 161-164 Advertisement
        {
            "text": """Questions 161-164 refer to the following advertisement.
Position: Assistant Editor | Date Posted: March 15
Description:
Goldhorse Press, an independent book publishing company focusing on North American gardening, has provided gardening advice for home gardeners for over 50 years. We are currently seeking an assistant editor to join our expanding team in Charlotte, North Carolina. We offer an excellent benefits package that includes medical and dental insurance.
Responsibilities:
Assist editors in the acquisition of titles with mainstream appeal; develop manuscripts, collaborate with outside support, and work closely with two managing editors.
Requirements/Qualifications:
• Bachelor's degree in a related field
• Minimum of one year of experience in the publishing industry
• Ability to pay close attention to detail
• Comfortable working as a team member
• Some experience in gardening preferred
• Immediate availability
E-mail resume and salary requirements to humanresources@goldhorsepress.com.""",
            "questions": [
                (161, "What title would most likely be published by Goldhorse Press?",
                 "Growing Your Baking Business", "A Tour Guide to North Carolina", "Planting Perennial Flowers", "The Efficient Executive",
                 "C", "Nhà xuất bản chuyên về làm vườn (gardening) nên 'Planting Perennial Flowers' là phù hợp nhất."),
                (162, "What is indicated about Goldhorse Press?",
                 "It is hiring two assistant editors.", "It provides insurance to employees.", "It publishes trade journals.", "It is a newly established company.",
                 "B", "Bài đăng tuyển nêu rõ quyền lợi: 'includes medical and dental insurance'."),
                (163, "What is a requirement for the advertised position?",
                 "A passion for gardening", "An aptitude for noticing details", "A master's degree in a related field", "An ability to work without supervision",
                 "B", "Yêu cầu: 'Ability to pay close attention to detail' đồng nghĩa với khả năng chú ý chi tiết."),
                (164, "What are applicants asked to do?",
                 "Submit a job application form", "Submit a list of references", "Indicate availability", "Indicate desired pay",
                 "D", "Yêu cầu ứng viên gửi: 'salary requirements' (mức lương mong muốn).")
            ]
        },
        # 165-168 Online chat session
        {
            "text": """Questions 165-168 refer to the following online chat session.
Derek Marshall [8:19 A.M.]
Hi, everyone. I wanted to give an update on our merger with Ridgewood, Inc., and see if you have had any meetings.

Mai Chung [8:20 A.M.]
Have we worked out which of our Derek Boutique locations will stay open?

Derek Marshall [8:21 A.M.]
Yes, Derek Boutique will keep 35 stores open with the full line of clothing. The other 12 will move their inventory to the Ridgewood, Inc., locations.

Nikita Tamboli [8:22 A.M.]
When should the moves be scheduled? This month?

Derek Marshall [8:23 A.M.]
No, the 12 locations don't need to be vacated until the end of next month. Could you schedule this for five weeks from now?

Anthony Rossi [8:24 A.M.]
I met over lunch with my managers earlier this week, including the two that I just hired.

Nikita Tamboli [8:25 A.M.]
That would work.

Mai Chung [8:27 A.M.]
I'll be meeting with the managers on my staff, too, to review the transition plan.

Derek Marshall [8:28 A.M.]
This is good work, everyone.

Anthony Rossi [8:29 A.M.]
When will the whole process be complete?

Derek Marshall [8:30 A.M.]
Within about 6 months.""",
            "questions": [
                (165, "Where do the writers most likely work?",
                 "At a marketing firm", "At a clothing company", "At a real estate agency", "At a newspaper publisher",
                 "B", "Họ làm việc tại chuỗi cửa hàng thời trang Derek Boutique ('full line of clothing')."),
                (166, "How many stores will be closing?",
                 "5", "6", "12", "35",
                 "C", "Derek nói giữ lại 35 cửa hàng và 12 cửa hàng còn lại sẽ đóng/chuyển hàng đi ('The other 12 will move their inventory')."),
                (167, "At 8:25 A.M., what does Ms. Tamboli most likely mean when she writes, 'That would work'?",
                 "The current inventory will be doubled.", "Ridgewood, Inc., will close in two weeks.", "Mr. Marshall will meet with the managers.", "She can schedule the moves in the proposed time frame.",
                 "D", "Đồng ý với mốc thời gian 5 tuần mà Derek đề xuất."),
                (168, "What is suggested about Ms. Chung and Mr. Rossi?",
                 "They supervise other employees.", "They often shop at Ridgewood, Inc.", "They recently went to lunch together.", "They have concerns about the merger.",
                 "A", "Cả hai đều có nhân sự quản lý cấp dưới thuộc quyền giám sát của mình ('my managers', 'managers on my staff').")
            ]
        },
        # 169-171 Article
        {
            "text": """Questions 169-171 refer to the following article.
From the Shadows to the Limelight
By Calum Ellwood
Dr. Esther Nujoma, an agricultural biotechnologist with the Namibia Institute of Applied Sciences, is the author of several books on the practical applications of biotechnology. ---[1]---. According to book critic Paige Kinnock of the London Daily Register, "Dr. Nujoma has increased the public's awareness of the role of biotechnology in daily life through her ability to translate highly complex scientific material into simple language."
Her latest work, Shining Behind Shadows, marks a departure from her usual subject matter. ---[2]---. Rather, the book highlights the lives and careers of twelve of her peers from Africa and Asia. The idea came to her three years ago at a conference in Chile. As Dr. Nujoma recalls, "Listening to one speaker after another, I realized that many of my colleagues are from regions of the world, such as Africa and Asia, that were underrepresented." ---[3]---.
The book is quite compelling. Dr. Nujoma brings to life the stories of how her peers were drawn to the profession and the effort they pour into their work.
The book does, however, fall short in one respect: it provides insights only into the lives and careers of those working in the field of agricultural biotechnology. I would have welcomed the stories of those specializing in animal, marine, or medical biotechnology, too. ---[4]---. Even so, Dr. Nujoma has again succeeded in creating a work that speaks to experts and laypeople alike.""",
            "questions": [
                (169, "What is NOT indicated about Dr. Nujoma?",
                 "She has visited Chile.", "She is a talented writer.", "She is based out of Namibia.", "She was interviewed by the London Daily Register.",
                 "D", "Nhà phê bình Paige Kinnock chỉ nhận xét về sách của cô, chứ không phỏng vấn cô."),
                (170, "What does Mr. Ellwood say about Shining Behind Shadows?",
                 "It focuses on scientists from Africa and Asia.", "It was released three years ago.", "It details why Dr. Nujoma chose her career.", "It describes various branches of biotechnology.",
                 "A", "Cuốn sách kể về cuộc đời 12 nhà khoa học đến từ châu Phi và châu Á."),
                (171, "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong? 'Its focus is not on applying biotechnology in real-life situations.'",
                 "[1]", "[2]", "[3]", "[4]",
                 "B", "Đứng trước câu 'Rather, the book highlights the lives and careers of twelve of her peers...'.")
            ]
        },
        # 172-175 Notice
        {
            "text": """Questions 172-175 refer to the following notice.
KINGSTON-GARNET ISLAND PASSENGER FERRY SERVICE
General Information
Beginning on 15 May, the Kingston-Garnet Island Passenger Ferry Service will resume service for eight weeks during the summer season. Ferries run daily every half hour from 7 a.m. to 8 p.m. The last ferry to Garnet Island will depart at 7:30 p.m. The last ferry from Garnet Island will leave at 8 p.m.
Bicycles are permitted on the passenger ferries. Bicyclists should arrive 30 minutes prior to departure and wait in the special bicycle lane to be loaded first. There are five racks that hold 50 bicycles on every passenger ferry.
No motorized vehicles are permitted on Garnet Island. Overnight parking is allowed in the main ferry terminal lot in Kingston. Rates are $5 per hour for up to 4 hours and a flat fee of $25 for four to 24 hours.
Visit our Web site at www.kgferryservice.com for photographs of the ferry boats, a map of Garnet Island, lists of local attractions on the island, and information about peak-hour fare increases and group discounts.""",
            "questions": [
                (172, "What is indicated about the ferry service?",
                 "It is available only seasonally.", "Its boats were recently upgraded.", "It takes an hour to reach the island.", "It runs more frequently on weekends.",
                 "A", "Dịch vụ chỉ chạy 8 tuần vào mùa hè ('available only seasonally')."),
                (173, "What is true about bicyclists on the ferries?",
                 "They are last to board the boat.", "They must purchase a special ticket.", "They cannot travel on the 7:30 P.M. trip.", "They should arrive at the terminal early.",
                 "D", "Người đi xe đạp cần có mặt trước giờ xuất bến 30 phút."),
                (174, "What is not allowed on Garnet Island?",
                 "Renting bicycles", "Driving cars", "Camping overnight", "Taking photographs",
                 "B", "Quy định cấm phương tiện cơ giới: 'No motorized vehicles are permitted on Garnet Island'."),
                (175, "What is indicated about the ferry tickets?",
                 "They can be purchased at stores in Kingston.", "They are less expensive for children.", "They vary in price depending on the time of travel.", "They are more expensive if purchased on the boat.",
                 "C", "Giá vé thay đổi theo thời gian cao điểm: 'information about peak-hour fare increases'.")
            ]
        },
        # 176-180 Double Passage (Receipt and E-mail)
        {
            "text": """Questions 176-180 refer to the following receipt and e-mail.
Receipt - Green Stripe Press
Order Number: GSP20896
Customer Information: Shoebox Mountain | Jason Ho <jasonho@shoeboxmountain.com>
Order Date: December 14 (PREPAID: online order)
Expected Delivery: December 18-20
Quantity | Item# | Item | Price
1 | CAL201 | Complimentary Wild Animals Calendar | $0.00
7 | ARB132 | Accounting Record Book ($19.99 each) | $139.93
Subtotal: $139.93
Tax @ 6%: $8.40
Shipping & Handling: $0.00
Total: $148.33
There is no charge for shipping and handling for corporate accounts. For questions regarding this order, please contact customerservice@greenstripepress.com.

---
E-mail:
To: customerservice@greenstripepress.com
From: Jason Ho <jasonho@shoeboxmountain.com>
Date: December 18
Subject: Order #GSP20896
Dear Green Stripe Press,
I am writing regarding my most recent order (#GSP20896), which was delivered today. If you check my original order, you will see that I ordered six accounting record books. You sent us seven copies. I'd like to return the one I didn't order and have our corporate credit card refunded, together with the shipping cost the return will incur. Let me know how you would like me to proceed.
On a different note, congratulations on the calendar you included in my order! The photos are even more stunning than those in the Ancient Castles calendar you sent us last year. Some staff members saw mine and want copies of their own. Would you mind sending two more our way?
Best,
Jason Ho, Owner, Shoebox Mountain""",
            "questions": [
                (176, "Why was Mr. Ho not charged a shipping fee?",
                 "He took advantage of a promotion.", "He made the purchase for his company.", "He picked up his order in person.", "He overpaid for shipping on a previous order.",
                 "B", "Hóa đơn ghi miễn phí vận chuyển cho tài khoản doanh nghiệp ('corporate accounts')."),
                (177, "What is true about Mr. Ho's order?",
                 "It was damaged in transit.", "It included fragile items.", "It was paid for by check.", "It was delivered on time.",
                 "D", "Hàng dự kiến đến ngày 18-20 tháng 12 và đã đến đúng ngày 18."),
                (178, "According to the e-mail, what was Mr. Ho sent by mistake?",
                 "A castle book he did not order", "A calendar for last year", "An incorrect refund check", "An extra accounting book",
                 "D", "Ông Ho đặt 6 cuốn nhưng bị gửi nhầm thành 7 cuốn ('An extra accounting book')."),
                (179, "What is one reason Mr. Ho wrote the e-mail?",
                 "To offer praise for an item", "To complain about a price", "To order some photo albums", "To recommend a graphic designer",
                 "A", "Ông gửi lời khen bộ lịch động vật hoang dã tặng kèm ('offer praise for an item')."),
                (180, "In the e-mail, the phrase 'our way' in paragraph 2, line 4, is closest in meaning to:",
                 "in our style", "to our address", "at our expense", "for our benefit",
                 "B", "'our way' trong ngữ cảnh gửi thêm 2 bản tới nghĩa là gửi về địa chỉ chúng tôi ('to our address').")
            ]
        },
        # 181-185 Double Passage (Contact Form and E-mail)
        {
            "text": """Questions 181-185 refer to the following contact form and e-mail.
Contact the Hilgrave Historical Commission
Name: Roger Witmondt
E-mail: roger@witmondtlocations.com
Comment:
I work as a movie location scout. At this time, I need to find a location for a client's short film. The film will be set in the early 1930s, and I understand that Hilgrave has some interesting architecture from that era.
Could you suggest places I should visit when I am in Hilgrave next month? Specifically, I am looking for a vacant building with an old-fashioned storefront that has distinctive details like stripe-patterned awnings and framed display windows. While having easy access to electricity would be ideal, it is not critical; my client has a generator that can be used if necessary. My client's aesthetic requirements are the most important consideration. Thank you for your assistance.

---
E-mail:
To: roger@witmondtlocations.com
From: brandischaertl@hilgravehistoricalcommission.org
Date: 28 September
Subject: The information you requested
Attachment: Hilgrave Sites
Dear Mr. Witmondt,
Thank you for reaching out to the Hilgrave Historical Commission. The attached brochure lists the main areas of historical interest in Hilgrave. Some buildings may not meet all your requirements, but they might be worth considering.
One building that is not on the attached list is the old record store at 188 Main Street. It has been closed to business for the past ten years or so, but it seems to meet your criteria. The owner, Luke Nylund, is currently using it for storage; I would be happy to put you in touch with him. I doubt you will have trouble getting permission from him to use the space. While Hilgrave was once a bustling town, these days there are fewer businesses in operation. Many of the town's residents would appreciate the attention a short film might bring to the area.
Best regards,
Brandi Schaertl""",
            "questions": [
                (181, "Why is Mr. Witmondt looking for a filming location in Hilgrave?",
                 "It is known for its scenic mountain views.", "It is a short distance from his office.", "Its business district has appeared in other films.", "Its buildings represent a particular time period.",
                 "D", "Phim lấy bối cảnh những năm 1930 và thị trấn có kiến trúc đại diện cho thời kỳ đó."),
                (182, "On the contact form, the word 'critical' in paragraph 2, line 4, is closest in meaning to:",
                 "judgmental", "essential", "sustainable", "available",
                 "B", "'not critical' nghĩa là không thiết yếu ('essential') vì đã có máy phát điện dự phòng."),
                (183, "What is implied about 188 Main Street?",
                 "It was once used as a residence.", "It is frequently visited by tourists.", "It has several floors.", "It has decorative design features.",
                 "D", "Tòa nhà đáp ứng các chi tiết cổ kính như mái hiên sọc và cửa kính trưng bày ('decorative design features')."),
                (184, "What does Ms. Schaertl offer to do?",
                 "Advertise a forthcoming film", "Arrange to have a building cleaned", "Connect Mr. Witmondt with a building's owner", "Help Mr. Witmondt acquire necessary permits from the town",
                 "C", "Đề nghị kết nối với ông chủ Luke Nylund ('put you in touch with him')."),
                (185, "What does Ms. Schaertl suggest that Hilgrave needs?",
                 "More publicity", "Additional parking", "A storage facility", "A business directory",
                 "A", "Thị trấn cần sự chú ý/quảng bá mà bộ phim mang lại ('more publicity').")
            ]
        },
        # 186-190 Triple Passage (E-mails and Web Page)
        {
            "text": """Questions 186-190 refer to the following e-mails and Web page.
E-mail 1:
To: Cornelia Payne <cpayne@roughwing.co.uk>
From: Pragya Mehta <pmehta@airsky.in>
Subject: Lecture | Date: 18 January
Attachment: Notes
Dear Cornelia,
I am sorry for the late notice, but I will not be able to join you for the lecture. My new position at the university in Mumbai requires me to remain on campus.
Although I know you are fully prepared to deliver the lecture on your own, I have attached a copy of the notes I had prepared for the presentation. After you review them, let me know if there is anything else I can add.
I had very much looked forward to traveling to Freeport, The Bahamas, for the first time and to seeing you again. I certainly miss working with you at the City University.
Good luck with the lecture.
Pragya

---
E-mail 2:
To: Ezra Halton <ehalton@ansonhouse.org>
From: Cornelia Payne <cpayne@roughwing.co.uk>
Subject: Information | Date: 20 January
Dear Mr. Halton,
I am pleased to confirm that 19 March works well for the lecture at the Anson House. As I had mentioned to you over the phone, I will be giving the lecture on my own.
My publisher, Alphagamma Press, will be sending 30 copies of my most recent work, coauthored with Dr. Pragya Mehta, to your institution. You and I will both receive an e-mail confirmation from Alphagamma when the books ship. They should arrive at least a week before the lecture.
I look forward to meeting you soon.
Cornelia Payne

---
Web page:
The Anson House, Freeport, The Bahamas - Speaker Series
Tickets for lectures are $35 each or may be purchased at a reduced rate of $90 for the series of three. Complimentary refreshments will be served. Reservations are not required but are recommended.
• 7 March, 7:00-8:30 P.M., Ms. Janelle Pears: Go behind the scenes of Anson House with historian Janelle Pears.
• 12 March, 7:00-8:30 P.M., Mr. Gregory Li: Restoring historic gardens.
• 19 March, 7:00-8:30 P.M., Dr. Cornelia Payne: Dr. Payne discusses her most recently published book, The Transatlantic World of the Nineteenth Century, coauthored with Dr. Pragya Mehta. Dr. Payne has taught at the City University of Stoke-on-Trent in England for nearly 25 years.""",
            "questions": [
                (186, "Why did Dr. Mehta send the e-mail to Dr. Payne?",
                 "To cancel a vacation plan", "To request lecture notes", "To confirm a meeting", "To offer an apology",
                 "D", "Gửi lời xin lỗi ('To offer an apology') vì không thể tham gia buổi thuyết trình."),
                (187, "According to the second e-mail, what did Dr. Payne do?",
                 "Arrange a delivery", "Give Mr. Halton a gift", "Revise her presentation", "Announce an address change",
                 "A", "Sắp xếp nhà xuất bản gửi 30 cuốn sách tới viện Anson House ('Arrange a delivery')."),
                (188, "What is indicated about Alphagamma Press?",
                 "It has headquarters in The Bahamas.", "It regularly ships materials to Anson House.", "It published The Transatlantic World of the Nineteenth Century.", "It is paying Dr. Payne's travel and accommodation expenses.",
                 "C", "Là nhà xuất bản cuốn sách The Transatlantic World of the Nineteenth Century của Dr. Payne và Dr. Mehta."),
                (189, "Where were Dr. Mehta and Dr. Payne most likely colleagues?",
                 "In Mumbai", "In Freeport", "In Stoke-on-Trent", "In Charleston",
                 "C", "Họ từng làm đồng nghiệp tại City University of Stoke-on-Trent."),
                (190, "What is indicated about the Speaker Series in the Web page?",
                 "Discounts are unavailable.", "Reservations are optional.", "Events are held in the morning.", "Refreshments are not included.",
                 "B", "Đặt chỗ là tùy chọn không bắt buộc ('Reservations are not required but are recommended').")
            ]
        },
        # 191-195 Triple Passage (Advertisement, E-mail, Press Release)
        {
            "text": """Questions 191-195 refer to the following advertisement, e-mail, and press release.
Advertisement:
CRYN Group: We find the best employees for your company.
Posted: November 25
Position title and codes:
• Director of Operations, TL0015
• Marketing Director, TL0023
• Quality Control Director, TL0027
• Director of Category Management, TL0045
To apply: Send your resume to mdoro@cryngroup.ca.com with posting number 2098 in the subject line.
Company: Our client plans to chart an aggressive growth path in Latin America, where it plans to start operations next year. It is a well-established company widely known for marketing consumer-health products in North America, and more recently in Europe and Asia.
Qualifications: Candidates must have a formal business degree and a proven management record in at least one international setting. Experience with online sales and marketing preferred.

---
E-mail:
To: Sven Arvidson <sarvidson@barkent.de.com>
From: Maria Doro <mdoro@cryngroup.ca.com>
Re: Posting number 2098 | Date: December 10
Dear Mr. Arvidson,
Thank you for submitting your resume. I'd like to schedule a preliminary telephone conversation with you as soon as possible to determine whether you would be a good choice for our client. I would particularly like to discuss whether you would be prepared to work at our client's new facility overseas.
Let me know if you are available for a 30-minute phone call sometime between 10:00 A.M. and 2:00 P.M. EST on Monday or Tuesday of next week. Please respond to me by e-mail at your earliest convenience.
Maria Doro, CRYN Group

---
Press Release:
For immediate release | Contact Julie Dryden at 416-555-0103
Tayerson Ltd. Names New Director
Toronto, April 5- Tayerson Ltd. continues to staff its new international outpost, which opens one month from today. In this case, the lucky candidate is Sven Arvidson. "Category Management is a relatively new area- and one that I'm sure I'll enjoy exploring and developing as director," Mr. Arvidson said. "As a major online marketer of nutritional supplements, Tayerson is poised to lead the way toward better living while becoming a more vibrant, more profitable corporation."
Mr. Arvidson has held several key management positions throughout his career, most recently at Barkent Pharmaceuticals in Germany. He was also a partner at MSZ Consulting Group, where he provided marketing guidance to leading consumer-product companies in Canada and China.""",
            "questions": [
                (191, "According to the advertisement, what qualification is not required of applicants?",
                 "A degree in business", "Previous employment in a managerial position", "Experience working overseas", "Online sales and marketing experience",
                 "D", "Kinh nghiệm bán hàng online là 'preferred' (ưu tiên) chứ không bắt buộc ('not required')."),
                (192, "What most likely is the location of the position for which Mr. Arvidson applied?",
                 "Latin America", "North America", "Europe", "Asia",
                 "A", "Vị trí tuyển dụng phục vụ mở rộng thị trường tại khu vực Mỹ Latinh ('Latin America')."),
                (193, "What does Ms. Doro ask Mr. Arvidson to provide in his reply?",
                 "Information on his leadership skills", "His availability for an interview", "A current resume", "The names of two references",
                 "B", "Hỏi thời gian có thể tham gia phỏng vấn qua điện thoại ('availability for an interview')."),
                (194, "What job code did Mr. Arvidson most likely reference in his application?",
                 "TL0015", "TL0023", "TL0027", "TL0045",
                 "D", "Vị trí của anh ấy là 'Director of Category Management' tương ứng mã TL0045."),
                (195, "According to the press release, what does Tayerson Ltd. sell?",
                 "Medical devices", "Accounting software", "Nutritional supplements", "Exercise equipment",
                 "C", "Bài báo ghi: 'major online marketer of nutritional supplements' (thực phẩm bổ sung dinh dưỡng).")
            ]
        },
        # 196-200 Triple Passage (Brochure, E-mail, Schedule)
        {
            "text": """Questions 196-200 refer to the following brochure, e-mail, and schedule.
Brochure:
FGJ's Business Expert Series: Prospect Data Purchasing
45-Minute Webinar | 11 May, 2:30 P.M.
Every department in an organisation needs information to make business decisions. Marketing professionals in particular rely on accurate data about potential customers to be successful. In this Webinar, Briana Carrera, chief financial officer of Pile One Market Data, explains how to avoid the pitfalls of acquiring incomplete data and provides tips on what to ask your prospective data provider to ensure that your next batch of data leads to the results you want.

---
E-mail:
To: Sandra Lescure
From: Gino Stelleti
Date: 12 May
Subject: FGJ's Webinar
Dear Sandra,
You had asked me to give you a report on the Webinar I attended yesterday about acquiring data. To be honest, even though the facilitator was clearly knowledgeable, she did not tell me anything I didn't already know. I was hoping that she would show us the differences between the many types of databases out there, but she never touched on that topic. I am not so sure all these Webinars are useful; in the end, they always seem to be steering us to purchase from one particular company or another-in this case it was Pile One. Having said that, the new schedule just came out, and there is another Webinar coming up soon that I am hoping will provide the information I want about how to store and sort data. I have already signed up for it; I will keep you posted about what I learn.
Gino

---
Schedule:
Upcoming 45-Minute Webinars in FGJ's Business Expert Series
• Basics of Market Research | 1 June, 9:30 A.M. | Ed Quinones | £12.00
• Tips for Accelerating Sales | 12 June, 9:30 A.M. | Cameron Stone | £16.00
• What is Market Automation? | 5 July, 11:00 A.M. | Ed Quinones | £12.00
• Choosing the Right Database | 17 July, 2:00 P.M. | Selina Tucci | £12.00
The Business Expert Series is organised by Tuyet Nguyen. Please address any questions to her at tnguyen@fgj.org. To access Webinar content, attendance is required as Webinars are not recorded for later viewing.""",
            "questions": [
                (196, "According to the brochure, who would benefit most from the Prospect Data Purchasing Webinar?",
                 "A financial analyst", "A marketing manager", "A customer-service representative", "An information-technology specialist",
                 "B", "Tài liệu ghi: 'Marketing professionals in particular rely on accurate data about potential customers'."),
                (197, "What was Mr. Stelleti's complaint about the Webinar?",
                 "He did not like the format.", "He did not learn anything new.", "He had difficulty hearing everything.", "He found the topic to be too complex.",
                 "B", "Gino viết: 'she did not tell me anything I didn't already know' (không học được điều gì mới)."),
                (198, "What does Mr. Stelleti suggest about Ms. Carrera?",
                 "She recently joined a new company.", "She has experience managing databases.", "She had asked him to sign up for the Webinar.", "She tried to sell her company's services.",
                 "D", "Gino nhận xét: 'seem to be steering us to purchase from one particular company or another-in this case it was Pile One'."),
                (199, "When will Mr. Stelleti most likely attend another Webinar?",
                 "On June 1", "On June 12", "On July 5", "On July 17",
                 "D", "Webinar về lưu trữ và sắp xếp dữ liệu ('how to store and sort data') là 'Choosing the Right Database' vào ngày 17 tháng 7."),
                (200, "According to the schedule, what do all of the Webinars have in common?",
                 "They cost the same amount.", "They are held in the morning.", "They last the same amount of time.", "They are recorded for future playback.",
                 "C", "Tất cả các buổi Webinar đều kéo dài 45 phút ('Upcoming 45-Minute Webinars').")
            ]
        }
    ]
    
    for p in part7_passages:
        doc.add_paragraph("[PASSAGE_START]")
        doc.add_paragraph(p["text"])
        doc.add_paragraph("[PASSAGE_END]")
        doc.add_paragraph("")
        
        for q in p["questions"]:
            q_num, text, opt_a, opt_b, opt_c, opt_d, ans, hint = q
            doc.add_paragraph(f"[Q: {q_num}] {text}")
            doc.add_paragraph("[SHUFFLE_Q: TRUE]")
            doc.add_paragraph("[SHUFFLE_A: TRUE]")
            doc.add_paragraph(f"[A] {opt_a}")
            doc.add_paragraph(f"[B] {opt_b}")
            doc.add_paragraph(f"[C] {opt_c}")
            doc.add_paragraph(f"[D] {opt_d}")
            doc.add_paragraph(f"[ANS: {ans}]")
            doc.add_paragraph(f"[HINT: {hint}]")
            doc.add_paragraph("[Q_END]")
            doc.add_paragraph("")
            
    doc.save(output_path)
    print(f"Generated: {output_path} successfully!")

def create_englishtest_docx(output_path, is_valid=True):
    doc = docx.Document()
    doc.add_paragraph("[EXAM_TITLE: TOEIC Reading Test 01 - EnglishTest Official]")
    doc.add_paragraph("[EXAM_DURATION: 75]")
    doc.add_paragraph("")
    doc.add_paragraph("[PART: 5]")
    
    data_101_130 = [
        (101, "Ms. Durkin asked for volunteers to help ------- with the employee fitness program.",
         "she", "her", "hers", "herself", "B", "Cần đại từ nhân xưng làm tân ngữ 'her' đứng sau động từ 'help' (giúp đỡ cô ấy)."),
        (102, "Lasner Electronics' staff have extensive ------- of current hardware systems.",
         "know", "known", "knowledge", "knowledgeable", "C", "Sau tính từ 'extensive' và trước giới từ 'of' cần một danh từ: 'knowledge' (kiến thức sâu rộng)."),
        (103, "------- a year, Tarrin Industrial Supply audits the accounts of all of its factories.",
         "Once", "Immediately", "Directly", "Yet", "A", "Cụm chỉ tần suất 'Once a year' có nghĩa là mỗi năm một lần."),
        (104, "Ms. Pham requested a refund ------- the coffeemaker she received was damaged.",
         "despite", "why", "concerning", "because", "D", "Cần liên từ chỉ nguyên nhân 'because' đứng trước mệnh đề 'the coffeemaker she received was damaged'."),
        (105, "Information ------- the artwork in the lobby is available at the reception desk.",
         "across", "about", "upon", "except", "B", "Cụm 'Information about something' mang nghĩa thông tin về cái gì đó."),
        (106, "With the Gema XTI binoculars, users can ------- see objects that are more than 100 meters away.",
         "ease", "easy", "easily", "easier", "C", "Cần trạng từ 'easily' đứng giữa động từ khuyết thiếu 'can' và động từ thường 'see' để bổ nghĩa cho hành động nhìn."),
        (107, "The Physical Therapy Association is committed to keeping costs ------- for its certification programs.",
         "affordable", "permitted", "cutting", "necessary", "A", "Cấu trúc 'keep + O + Adj' -> 'keeping costs affordable' (giữ chi phí ở mức hợp lý/phải chăng)."),
        (108, "Mr. Brennel ------- positions in various areas of the company before he became president.",
         "occupation", "occupational", "occupying", "occupied", "D", "Câu thiếu động từ chính chia ở thì quá khứ đơn ('occupied' - đã nắm giữ các vị trí) trước mệnh đề trạng ngữ chỉ thời gian 'before he became president'."),
        (109, "To remain on schedule, editors must submit all ------- to the book to the authors by Friday.",
         "ideas", "essays", "revisions", "suggestions", "C", "'Revisions to the book' có nghĩa là các nội dung chỉnh sửa/sửa đổi đối với cuốn sách."),
        (110, "------- industry professionals are allowed to purchase tickets to the Kuo Photography Fair.",
         "Only", "Until", "Unless", "Quite", "A", "Trạng từ 'Only' đứng đầu câu bổ nghĩa cho chủ ngữ: 'Chỉ những chuyên gia trong ngành mới được phép mua vé...'."),
        (111, "At Pharmbeck's banquet, Mr. Jones ------- a trophy for his performance in this year's quality-improvement initiative.",
         "accepted", "congratulated", "nominated", "hoped", "A", "'Accepted a trophy' có nghĩa là đã đón nhận chiếc cúp khen thưởng."),
        (112, "Ms. Suto claims that important market trends become ------- with the use of data analysis.",
         "predict", "prediction", "predictable", "predictably", "C", "Sau động từ nối 'become' cần một tính từ: 'predictable' (có thể dự đoán được)."),
        (113, "One of Grommer Consulting's goals is to enhance the relationship ------- salespeople and their customers.",
         "inside", "within", "around", "between", "D", "Cấu trúc 'relationship between A and B' (mối quan hệ giữa nhân viên bán hàng và khách hàng của họ)."),
        (114, "Depending on your answers to the survey, we ------- you to collect additional information.",
         "may call", "are calling", "have been called", "must be calling", "A", "'may call' diễn tả một khả năng có thể xảy ra trong tương lai tùy thuộc vào kết quả khảo sát."),
        (115, "------- Jemburger opened its newest franchise, the first 100 customers were given free hamburgers.",
         "Now", "When", "As if", "After all", "B", "Cần liên từ chỉ thời gian 'When' nối hai mệnh đề quá khứ đơn (Khi Jemburger mở cửa hàng nhượng quyền mới nhất...)."),
        (116, "Please include the serial number of your product in any ------- with the customer service department.",
         "corresponds", "correspondence", "correspondingly", "correspondent", "B", "Sau lượng từ 'any' cần danh từ 'correspondence' (thư từ/liên lạc giao dịch với bộ phận CSKH)."),
        (117, "The award-winning film Underwater Secrets promotes awareness ------- ocean pollution and its effects on our planet.",
         "of", "to", "from", "with", "A", "Cụm cố định 'promote awareness of something' (nâng cao nhận thức về vấn đề gì)."),
        (118, "BYF Company specializes in ------- promotional items to help companies advertise their brand.",
         "personally", "personalized", "personality", "personalizes", "B", "Cần tính từ 'personalized' (được cá nhân hóa) đứng trước bổ nghĩa cho cụm danh từ 'promotional items'."),
        (119, "------- the rent increase is less than 2 percent, Selwin Electrical Supply will continue to lease the space.",
         "As long as", "Along with", "Not only", "Otherwise", "A", "Liên từ điều kiện 'As long as' mang nghĩa 'Miễn là' (Miễn là mức tăng tiền thuê dưới 2%...)."),
        (120, "Belden Hospital's chief of staff meets regularly with the staff to ensure that procedures ------- correctly.",
         "to be performed", "would have performed", "had been performed", "are being performed", "D", "Mệnh đề sau 'that' cần động từ chia ở thể bị động hiện tại 'are being performed' để hòa hợp về thì với 'meets regularly'."),
        (121, "Any requests for time off should be addressed to the ------- department supervisor.",
         "urgent", "appropriate", "subsequent", "deliverable", "B", "'the appropriate department supervisor' nghĩa là người giám sát bộ phận phù hợp/tương ứng."),
        (122, "World Fish Supply delivers the freshest fish possible thanks to innovative ------- and shipping methods.",
         "preserves", "preserved", "preserve", "preservation", "D", "Cần danh từ 'preservation' (sự bảo quản) song hành với 'shipping' để tạo cụm 'preservation and shipping methods'."),
        (123, "Company executives are currently reviewing the annual budget ------- submitted to them by the Financial Planning department.",
         "requirements", "deliveries", "developers", "qualities", "A", "'annual budget requirements' có nghĩa là các yêu cầu/đề xuất ngân sách hàng năm."),
        (124, "Even the CEO had to admit that Prasma Designs' win was ------- the result of fortunate timing.",
         "parts", "parted", "partly", "parting", "C", "Trạng từ 'partly' (một phần) bổ nghĩa cho cụm 'the result of fortunate timing'."),
        (125, "Mr. Singh took notes on ------- the focus group discussed during the morning session.",
         "each", "several", "another", "everything", "D", "Đại từ 'everything' đóng vai trò tân ngữ của giới từ 'on' và được bổ nghĩa bởi mệnh đề 'the focus group discussed' (mọi thứ mà nhóm đã thảo luận)."),
        (126, "Last year, Tadaka Computer Solutions ranked third ------- in regional earnings.",
         "together", "overall", "consecutively", "generally", "B", "Cụm 'ranked third overall' có nghĩa là xếp hạng ba tính trên tổng thể."),
        (127, "------- the popularity of the BPT39 wireless speaker, production will be increased fivefold starting next month.",
         "On behalf of", "Whether", "Moreover", "As a result of", "D", "Cụm giới từ chỉ nguyên nhân 'As a result of + Noun phrase' (Do kết quả của sự phổ biến...)."),
        (128, "Zypo Properties has just signed a lease agreement with the law firm ------- offices are on the third floor.",
         "how", "what", "whose", "wherever", "C", "Đại từ quan hệ sở hữu 'whose' đứng trước danh từ 'offices' (công ty luật mà văn phòng của họ ở tầng 3)."),
        (129, "------- events this year caused profits in the second and third quarters to differ significantly from original projections.",
         "Total", "Marginal", "Representative", "Unforeseen", "D", "Tính từ 'Unforeseen' (không lường trước được) bổ nghĩa cho 'events' giải thích lý do lợi nhuận khác biệt lớn so với dự kiến."),
        (130, "The timeline for the pathway lighting project was extended to ------- input from the environmental commission.",
         "use up", "believe in", "make into", "allow for", "D", "Cụm động từ 'allow for' mang nghĩa tính đến / cho phép có thêm thời gian để tiếp nhận ý kiến đóng góp.")
    ]
    
    for q in data_101_130:
        q_num, text, opt_a, opt_b, opt_c, opt_d, ans, hint = q
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: TRUE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        doc.add_paragraph(f"[A] {opt_a}")
        doc.add_paragraph(f"[B] {opt_b}")
        doc.add_paragraph(f"[C] {opt_c}")
        if not is_valid and q_num == 105:
            pass # omit [D] for test 1
        else:
            doc.add_paragraph(f"[D] {opt_d}")
        doc.add_paragraph(f"[ANS: {ans}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[Q_END]")
        doc.add_paragraph("")
        
    # Part 6
    doc.add_paragraph("[PART: 6]")
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("""Questions 131-134 refer to the following advertisement.
With Global Strength Gym's 30-day trial period, you get the opportunity to try out our classes, equipment, and facilities. It's completely risk-free! To sign up, we require your contact information and payment details, but you will only be charged if you are a member for ---[132]--- 30 days. ---[131]---. If you decide within this time that you no longer want to be a member of Global Strength, visit our Web site at www.gsgym.com. On the Membership page, elect to ---[134]--- your membership and ---[133]--- enter the necessary information. It's that easy!""")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("")
    
    p6_data = [
        (131, "Chọn câu phù hợp điền vào vị trí [131]:",
         "Throughout the trial, you pay nothing and sign no contract.", "Weight-lifting classes are not currently available.", "A cash deposit is required when you sign up for membership.", "All questions should be e-mailed to customerservice@gsgym.com.",
         "A", "Đoạn văn đang nói về gói dùng thử 30 ngày miễn phí rủi ro ('completely risk-free'), nên câu A là phù hợp nhất."),
        (132, "Chọn từ phù hợp điền vào vị trí [132]:",
         "not even", "almost", "over", "less than",
         "C", "Bạn chỉ bị tính phí nếu là thành viên trên 30 ngày ('over 30 days')."),
        (133, "Chọn từ phù hợp điền vào vị trí [133]:",
         "justly", "regularly", "evenly", "simply",
         "D", "Trạng từ 'simply' (chỉ cần đơn giản là nhập thông tin cần thiết), khớp với câu kết 'It\\'s that easy!'."),
        (134, "Chọn từ phù hợp điền vào vị trí [134]:",
         "extend", "renew", "cancel", "initiate",
         "C", "Nếu không muốn làm thành viên nữa ('no longer want to be a member'), hãy chọn hủy ('cancel') tư cách thành viên.")
    ]
    
    for q in p6_data:
        q_num, text, opt_a, opt_b, opt_c, opt_d, ans, hint = q
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        doc.add_paragraph(f"[A] {opt_a}")
        doc.add_paragraph(f"[B] {opt_b}")
        doc.add_paragraph(f"[C] {opt_c}")
        doc.add_paragraph(f"[D] {opt_d}")
        doc.add_paragraph(f"[ANS: {ans}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[Q_END]")
        doc.add_paragraph("")
        
    doc.save(output_path)
    print(f"Generated: {output_path} successfully!")

if __name__ == "__main__":
    os.makedirs("sample_exams", exist_ok=True)
    
    # 1. EnglishTest: Part 5 (101-130) + Part 6 (131-134)
    create_englishtest_docx("sample_exams/EnglishTest_Part5_Full_Tagged.docx", is_valid=True)
    
    # 2. ETS Reading Test 09: Part 5, 6, 7 (101-200)
    create_sample_docx("sample_exams/ETS_Reading_Test09_Full_Tagged.docx", is_valid=True)
    
    # 3. Invalid file (missing [D] at Q105) for testing All-or-Nothing rollback
    create_englishtest_docx("sample_exams/EnglishTest_Invalid_MissingD.docx", is_valid=False)

