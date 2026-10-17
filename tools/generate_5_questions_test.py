import docx
from docx.shared import Pt, Inches, RGBColor

def create_test_5_questions_docx(output_path):
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = [
        "[EXAM_TITLE: TOEIC Reading Quick Test - 5 Câu Mẫu]",
        "[EXAM_DURATION: 10]",
        "",
        "[PART: 5]",
        "",
        # Câu 101
        "[Q: 101] Ms. Tanaka will present the quarterly budget report ------- the board meeting tomorrow.",
        "[SHUFFLE_Q: TRUE]",
        "[SHUFFLE_A: TRUE]",
        "[A] during",
        "[B] while",
        "[C] since",
        "[D] until",
        "[ANS: A]",
        "[HINT: \"during\" đi kèm danh từ chỉ sự kiện \"the board meeting\" (trong suốt cuộc họp). \"while\" cần một mệnh đề đầy đủ (S + V).]",
        "[Q_END]",
        "",
        # Câu 102
        "[Q: 102] All employees are encouraged to ------- in the upcoming leadership training seminar.",
        "[SHUFFLE_Q: TRUE]",
        "[SHUFFLE_A: TRUE]",
        "[A] participation",
        "[B] participate",
        "[C] participant",
        "[D] participating",
        "[ANS: B]",
        "[HINT: Cấu trúc \"encourage someone to + V nguyên mẫu\". Lựa chọn (B) participate là động từ nguyên mẫu đi kèm giới từ in.]",
        "[Q_END]",
        "",
        # Câu 103
        "[Q: 103] The new software update will significantly improve system ------- and reduce loading times.",
        "[SHUFFLE_Q: TRUE]",
        "[SHUFFLE_A: TRUE]",
        "[A] perform",
        "[B] performed",
        "[C] performance",
        "[D] performing",
        "[ANS: C]",
        "[HINT: Sau danh từ ghép \"system\" cần một danh từ để tạo thành cụm danh từ \"system performance\" (hiệu năng hệ thống).]",
        "[Q_END]",
        "",
        # Câu 104
        "[Q: 104] Please review the attached contract carefully ------- signing and returning it to HR.",
        "[SHUFFLE_Q: TRUE]",
        "[SHUFFLE_A: TRUE]",
        "[A] before",
        "[B] prior",
        "[C] ahead",
        "[D] forward",
        "[ANS: A]",
        "[HINT: Giới từ \"before\" có thể kết hợp trực tiếp với V-ing (\"before signing\" = trước khi ký). Lựa chọn (B) prior phải đi với \"to\".]",
        "[Q_END]",
        "",
        # Câu 105
        "[Q: 105] Customers who experience technical difficulties should contact customer support ------- .",
        "[SHUFFLE_Q: TRUE]",
        "[SHUFFLE_A: TRUE]",
        "[A] immediate",
        "[B] immediately",
        "[C] immediacy",
        "[D] immediateness",
        "[ANS: B]",
        "[HINT: Đứng cuối câu bổ nghĩa cho cụm động từ hành động \"contact customer support\", cần dùng trạng từ chỉ cách thức: \"immediately\" (ngay lập tức).]",
        "[Q_END]"
    ]

    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        if line.startswith("[EXAM_TITLE"):
            run.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 51, 102)
        elif line.startswith("[PART"):
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(30, 80, 160)
        elif line.startswith("[Q:"):
            run.bold = True
        elif line.startswith("[ANS:"):
            run.font.color.rgb = RGBColor(0, 128, 0)
        elif line.startswith("[HINT:"):
            run.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(output_path)
    print(f"Created docx: {output_path}")

if __name__ == "__main__":
    create_test_5_questions_docx(r"d:\Flutter VS Code Project\KT\flutter_application_1\Test_5_Questions_Sample.docx")
    create_test_5_questions_docx(r"d:\Flutter VS Code Project\KT\flutter_application_1\sample_exams\Test_5_Questions_Sample.docx")
