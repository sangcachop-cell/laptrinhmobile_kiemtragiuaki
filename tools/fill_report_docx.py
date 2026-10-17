import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import shutil
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def fill_report(template_path, output_path):
    print(f"Loading template from {template_path}...")
    doc = docx.Document(template_path)

    # 1. Điền thông tin sinh viên vào P0, P1, P2 (giữ nguyên định dạng màu và font của template)
    # P0: Họ và tên:
    p0 = doc.paragraphs[0]
    r0 = p0.add_run(" Nguyễn Tấn Sang")
    r0.font.name = 'Times New Roman'
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0x15, 0x3D, 0x63)

    # P1: Mã sinh viên:
    p1 = doc.paragraphs[1]
    r1 = p1.add_run(" 2380601899")
    r1.font.name = 'Times New Roman'
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x15, 0x3D, 0x63)

    # P2: Lớp:
    p2 = doc.paragraphs[2]
    r2 = p2.add_run(" 23DTHB2")
    r2.font.name = 'Times New Roman'
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x15, 0x3D, 0x63)

    # Định nghĩa nội dung mô tả ngắn gọn và ảnh chụp cho từng mục
    # Thứ tự các section trong template:
    # Item 1: Sơ đồ database (sau tiêu đề tại P5)
    # Item 2: Hình ảnh giao diện thiết kế trên Figma (sau tiêu đề tại P7)
    # Item 3: Giao diện Admin: Import đề thi (sau tiêu đề tại P9)
    # Item 4: Giao diện Admin: Chỉnh sửa nội dung đề thi (sau tiêu đề tại P11)
    # Item 5: Giao diện Users: Chọn lựa đề thi (sau tiêu đề tại P13)
    # Item 6: Giao diện Users: Làm bài thi (sau tiêu đề tại P15)
    # Item 7: Giao diện Users: Hiển thị kết quả bài thi (sau tiêu đề tại P17)

    sections_data = [
        {
            "header_keyword": "Sơ đồ database",
            "desc": "Mô tả ngắn gọn: Cơ sở dữ liệu Microsoft SQL Server (database: 2380601899Db trên máy chủ (localdb)\\mssqllocaldb) được thiết kế theo chuẩn 3NF với 6 bảng thực thể liên kết chặt chẽ: Exams (thông tin đề thi), Passages (đoạn văn bài đọc hiểu Part 6, 7), Questions (ngân hàng câu hỏi), QuestionOptions (4 lựa chọn A, B, C, D), ExamAttempts (lượt làm bài thi của thí sinh) và AttemptAnswers (chi tiết đáp án từng câu). Toàn bộ quan hệ cha - con đều được thiết lập khóa ngoại Foreign Key với ràng buộc toàn vẹn ON DELETE CASCADE bảo đảm an toàn dữ liệu.",
            "image": "docs/images/erd_schema.jpg",
            "width": 6.2
        },
        {
            "header_keyword": "Hình ảnh giao diện thiết kế trên Figma",
            "desc": "Mô tả ngắn gọn: Thiết kế giao diện chuẩn Mobile (375x812 px) trên Figma với triết lý tối giản, hiện đại và tập trung cao. Ứng dụng hỗ trợ cơ chế vào thi ngay không cần đăng nhập (No-Auth), phòng thi phân chia 2 nửa (Split-View 45% bài đọc / 55% câu hỏi), đồng hồ bấm giờ, Bảng câu hỏi chuyển nhanh (Palette) và chức năng Nộp bài sớm chỉ sau 1-2 câu. Đường link xem và tương tác trực tiếp Figma Prototype: docs/figma_interactive_prototype.html.",
            "image": "docs/images/figma_screens_flow.jpg",
            "width": 6.2
        },
        {
            "header_keyword": "Giao diện Admin: Import đề thi",
            "desc": "Mô tả ngắn gọn: Cổng quản trị Admin tại địa chỉ /Exam cho phép giảng viên tải lên tệp Word (.docx) đề thi TOEIC Reading chuẩn hóa. Hệ thống backend ASP.NET Core tự động bóc tách đoạn văn bài đọc, câu hỏi, các lựa chọn A/B/C/D và đáp án chính xác (tự động nhận diện từ gạch chân), tạo đề thi mới và kích hoạt ca thi công khai ngay lập tức.",
            "image": "docs/images/admin_upload_list.png",
            "width": 6.0
        },
        {
            "header_keyword": "Giao diện Admin: Chỉnh sửa nội dung đề thi",
            "desc": "Mô tả ngắn gọn: Giao diện xem trước và chỉnh sửa Admin Preview & Edit tại /Exam/Preview/1 cho phép xem lại toàn bộ cấu trúc bài thi vừa bóc tách, chỉnh sửa nội dung bài đọc, câu hỏi, thay đổi đáp án đúng, bổ sung gợi ý giải thích (Hint) và cấu hình trộn đề / trộn đáp án ngẫu nhiên.",
            "image": "docs/images/admin_preview_edit.png",
            "width": 6.0
        },
        {
            "header_keyword": "Giao diện Users: Chọn lựa đề thi",
            "desc": "Mô tả ngắn gọn: Trên ứng dụng Flutter di động, thí sinh mở app là thấy ngay danh sách các đề thi đã phát hành mà không cần qua bất kỳ bước đăng ký hay đăng nhập tài khoản nào. Thẻ đề thi hiển thị trực quan tiêu đề bài thi, thời gian làm bài (75 phút), số lượng câu hỏi và nút 'Vào thi ngay'.",
            "image": "docs/images/user_exam_list.jpg",
            "width": 3.6
        },
        {
            "header_keyword": "Giao diện Users: Làm bài thi",
            "desc": "Mô tả ngắn gọn: Giao diện phòng thi chia màn hình Split-View thông minh: nửa trên cuộn đọc tài liệu bài đọc hiểu giữ nguyên các vị trí điền khuyết [131], [132]..., nửa dưới chọn đáp án A, B, C, D. Đồng hồ đếm ngược hiển thị liên tục, hỗ trợ Bảng câu hỏi (Question Palette) và nút 'Nộp bài sớm' ngay khi mới hoàn thành 1-2 câu.",
            "image": "docs/images/user_exam_room.jpg",
            "width": 3.6
        },
        {
            "header_keyword": "Giao diện Users: Hiển thị kết quả bài thi",
            "desc": "Mô tả ngắn gọn: Màn hình kết quả ngay sau khi nộp bài: Hiển thị điểm số thang 10 và điểm quy đổi TOEIC Reading chuẩn thang 495, thống kê 3 màu số câu Đúng / Sai / Bỏ qua, cùng danh sách xem lại toàn bộ bài làm kèm lời giải thích chi tiết (Hint) phân tích ngữ pháp cho từng câu hỏi.",
            "image": "docs/images/user_exam_result.jpg",
            "width": 3.6
        }
    ]

    # Tìm vị trí các header trong document
    for sec in sections_data:
        kw = sec["header_keyword"]
        target_idx = None
        for i, p in enumerate(doc.paragraphs):
            if kw.lower() in p.text.lower():
                target_idx = i
                break
        
        if target_idx is not None:
            # Đoạn văn ngay sau header là target_idx + 1
            content_p = doc.paragraphs[target_idx + 1]
            content_p.text = "" # Xóa rỗng nếu có khoảng trắng thừa
            
            # Gán text mô tả với font Times New Roman, cỡ chữ 12pt, màu đen chuẩn
            r_desc = content_p.add_run(sec["desc"])
            r_desc.font.name = 'Times New Roman'
            r_desc.font.size = Pt(12)
            r_desc.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            content_p.paragraph_format.line_spacing = 1.15
            content_p.paragraph_format.space_after = Pt(6)

            # Chèn thêm đoạn văn chứa hình ảnh ngay sau content_p
            # Để chèn ngay sau content_p, ta tìm đoạn văn tiếp theo và insert_paragraph_before
            if target_idx + 2 < len(doc.paragraphs):
                next_p = doc.paragraphs[target_idx + 2]
                img_p = next_p.insert_paragraph_before()
            else:
                img_p = doc.add_paragraph()

            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_before = Pt(4)
            img_p.paragraph_format.space_after = Pt(12)

            img_path = sec["image"]
            if os.path.exists(img_path):
                r_img = img_p.add_run()
                r_img.add_picture(img_path, width=Inches(sec["width"]))
                print(f"  -> Added image {img_path} for '{kw}'")
            else:
                print(f"  -> Warning: Image {img_path} not found!")

    # 3. Điền bảng Đánh giá (Table 0)
    print("Filling evaluation table (Table 0)...")
    if doc.tables:
        tbl = doc.tables[0]
        # Duyệt từ dòng 1 đến 8
        for row_idx in range(1, len(tbl.rows)):
            row = tbl.rows[row_idx]
            
            # Cột 0: Số thứ tự 1..8
            c0 = row.cells[0]
            c0.text = ""
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(str(row_idx))
            r0.font.name = 'Times New Roman'
            r0.font.size = Pt(12)
            r0.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

            # Cột 3: Đã thực hiện (đánh dấu x)
            c3 = row.cells[3]
            c3.text = ""
            p3 = c3.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r3 = p3.add_run("x")
            r3.font.name = 'Times New Roman'
            r3.font.size = Pt(12)
            r3.font.bold = True
            r3.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

        # Bổ sung link prototype vào dòng 6
        cell_link = tbl.rows[6].cells[1]
        p_link = cell_link.add_paragraph()
        p_link.paragraph_format.space_before = Pt(2)
        r_link = p_link.add_run("(Link: docs/figma_interactive_prototype.html)")
        r_link.font.name = 'Times New Roman'
        r_link.font.size = Pt(10)
        r_link.font.italic = True
        r_link.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    # Lưu file
    doc.save(output_path)
    print(f"Successfully saved to {output_path}!")

if __name__ == "__main__":
    template = "BaoCao-HovaTen-Masinhvien.docx.bak"
    out1 = "BaoCao-NguyenTanSang-2380601899.docx"
    out2 = "BaoCao-HovaTen-Masinhvien.docx"
    
    fill_report(template, out1)
    fill_report(template, out2)
